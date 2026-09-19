"""
Reports router — dashboard stats, downtime, utilization, audit log, export (XLSX + PDF).
Replaces reportController.js entirely. All data from PostgreSQL.
"""
import io
from datetime import date, datetime, timedelta
from typing import Optional
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, and_

from backend.db.database import get_db
from backend.core.deps import get_current_user
from backend.models.block_schedule import BlockSchedule
from backend.models.maintenance_task import MaintenanceTask
from backend.models.audit_log import AuditLog
from backend.models.corridor import CorridorBlock, BlockWindow
from backend.models.user import User

router = APIRouter(tags=["reports"])

DEPT_DEFAULTS = {"Engineering": 0, "Traction Distribution": 0, "Signal & Telecom": 0}


def _fmt_duration(mins: int) -> str:
    if not mins:
        return "0min"
    h, m = divmod(int(mins), 60)
    if h and m:
        return f"{h}hr {m}min"
    return f"{h}hr" if h else f"{m}min"


@router.get("/dashboard-stats")
def dashboard_stats(
    zone: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    today = date.today()

    # Determine zone corridors if zone is provided
    zone_filter = zone and zone.strip().upper() != "ALL"
    z = zone.strip() if zone_filter else None
    zone_corridor_ids = None
    if zone_filter:
        zone_corridors = db.query(CorridorBlock).filter(
            (CorridorBlock.zone_code.ilike(z)) | (CorridorBlock.zone.ilike(z))
        ).all()
        zone_corridor_ids = {c.section_id for c in zone_corridors}

    task_q = db.query(MaintenanceTask.status, func.count().label("cnt"))
    if zone_corridor_ids is not None:
        task_q = task_q.filter(
            (MaintenanceTask.zone_code.ilike(z)) | (MaintenanceTask.section_id.in_(zone_corridor_ids))
        )
    status_counts = dict(task_q.group_by(MaintenanceTask.status).all())
    total_tasks = sum(status_counts.values())
    pending_tasks = status_counts.get("pending", 0)

    sched_q = db.query(BlockSchedule.status, func.count().label("cnt"))
    if zone_corridor_ids is not None:
        sched_q = sched_q.filter(
            (BlockSchedule.zone_code.ilike(z)) | (BlockSchedule.section_id.in_(zone_corridor_ids))
        )
    sched_status_counts = dict(sched_q.group_by(BlockSchedule.status).all())
    scheduled_blocks = sched_status_counts.get("proposed", 0)
    approved_blocks = sched_status_counts.get("approved", 0)

    overdue_q = db.query(func.count(MaintenanceTask.id)).filter(
        MaintenanceTask.criticality.in_(["critical", "high"]),
        MaintenanceTask.due_date < today,
        MaintenanceTask.status.in_(["pending", "scheduled"]),
    )
    if zone_corridor_ids is not None:
        overdue_q = overdue_q.filter(
            (MaintenanceTask.zone_code.ilike(z)) | (MaintenanceTask.section_id.in_(zone_corridor_ids))
        )
    overdue_critical = overdue_q.scalar() or 0

    # Dept breakdown (single GROUP BY query)
    dept_q = db.query(
        MaintenanceTask.department, MaintenanceTask.status, func.count().label("cnt")
    )
    if zone_corridor_ids is not None:
        dept_q = dept_q.filter(
            (MaintenanceTask.zone_code.ilike(z)) | (MaintenanceTask.section_id.in_(zone_corridor_ids))
        )
    dept_rows = dept_q.group_by(MaintenanceTask.department, MaintenanceTask.status).all()

    dept_all = {**DEPT_DEFAULTS}
    dept_pending = {**DEPT_DEFAULTS}
    for dept, st, cnt in dept_rows:
        if dept in dept_all:
            dept_all[dept] = dept_all[dept] + cnt
            if st == "pending":
                dept_pending[dept] = dept_pending[dept] + cnt

    # Criticality distribution (single GROUP BY query)
    crit_q = db.query(
        MaintenanceTask.criticality, func.count().label("cnt")
    )
    if zone_corridor_ids is not None:
        crit_q = crit_q.filter(
            (MaintenanceTask.zone_code.ilike(z)) | (MaintenanceTask.section_id.in_(zone_corridor_ids))
        )
    crit_rows = crit_q.group_by(MaintenanceTask.criticality).all()
    crit_dist = {"critical": 0, "high": 0, "medium": 0, "low": 0}
    for crit, cnt in crit_rows:
        if crit in crit_dist:
            crit_dist[crit] = cnt

    # Asset availability — single aggregate query
    if zone_corridor_ids is not None:
        corridor_count = len(zone_corridor_ids) or 1
    else:
        corridor_count = db.query(func.count(CorridorBlock.section_id)).scalar() or 12
    weekly_capacity = corridor_count * 10080  # 7 days * 1440 minutes per corridor

    downtime_q = db.query(func.sum(BlockSchedule.total_duration_min)).filter(
        BlockSchedule.plan_type == "weekly",
        BlockSchedule.status.in_(["approved", "proposed", "executed"])
    )
    if zone_corridor_ids is not None:
        downtime_q = downtime_q.filter(
            (BlockSchedule.zone_code.ilike(z)) | (BlockSchedule.section_id.in_(zone_corridor_ids))
        )
    total_downtime = downtime_q.scalar() or 0
    availability = round(((weekly_capacity - total_downtime) / weekly_capacity) * 100, 1) if weekly_capacity > 0 else 99.5

    # 7-day availability trend — batch query instead of 7 individual queries
    week_ago = today - timedelta(days=6)
    week_start_dt = datetime.combine(week_ago, datetime.min.time())
    week_end_dt = datetime.combine(today, datetime.max.time())

    # Fetch all relevant blocks in one query and aggregate in Python
    trend_q = db.query(
        BlockSchedule.window_start,
        BlockSchedule.total_duration_min,
    ).filter(
        BlockSchedule.window_start >= week_start_dt,
        BlockSchedule.window_start <= week_end_dt,
        BlockSchedule.status.in_(["approved", "proposed"]),
    )
    if zone_corridor_ids is not None:
        trend_q = trend_q.filter(
            (BlockSchedule.zone_code.ilike(z)) | (BlockSchedule.section_id.in_(zone_corridor_ids))
        )
    recent_blocks = trend_q.all()

    # Group downtime by day
    daily_downtime: dict[date, int] = {}
    for ws, dur in recent_blocks:
        if ws:
            d = ws.date()
            daily_downtime[d] = daily_downtime.get(d, 0) + (dur or 0)

    trend = []
    for i in range(6, -1, -1):
        target_date = today - timedelta(days=i)
        day_str = target_date.strftime("%a")
        downtime = daily_downtime.get(target_date, 0)
        daily_capacity = corridor_count * 1440
        day_avail = max(0.0, round(((daily_capacity - downtime) / daily_capacity) * 100, 1)) if daily_capacity > 0 else 99.5
        trend.append({
            "date": target_date.isoformat(),
            "day": day_str,
            "availability": day_avail,
            "downtimeMinutes": downtime,
        })

    # Recent schedules (top 5 proposed or approved)
    recent_sched_q = db.query(BlockSchedule).filter(
        BlockSchedule.status.in_(["proposed", "approved"])
    )
    if zone_corridor_ids is not None:
        recent_sched_q = recent_sched_q.filter(
            (BlockSchedule.zone_code.ilike(z)) | (BlockSchedule.section_id.in_(zone_corridor_ids))
        )
    recent_scheds = recent_sched_q.order_by(
        BlockSchedule.window_start.asc()
    ).limit(5).all()

    # Availability by section
    downtime_agg_q = db.query(BlockSchedule.section_id, func.sum(BlockSchedule.total_duration_min)).filter(
        BlockSchedule.plan_type == "weekly",
        BlockSchedule.status.in_(["approved", "proposed"])
    )
    if zone_corridor_ids is not None:
        downtime_agg_q = downtime_agg_q.filter(
            (BlockSchedule.zone_code.ilike(z)) | (BlockSchedule.section_id.in_(zone_corridor_ids))
        )
    downtime_agg = dict(downtime_agg_q.group_by(BlockSchedule.section_id).all())

    sec_q = db.query(CorridorBlock)
    if zone_corridor_ids is not None:
        sec_q = sec_q.filter(CorridorBlock.section_id.in_(zone_corridor_ids))
    sections = sec_q.all()

    avail_by_section = []
    for sec in sections:
        sec_window = 10080  # 7 days * 1440 min operating capacity per section
        sec_downtime = downtime_agg.get(sec.section_id) or 0
        avail_by_section.append({
            "sectionId": sec.section_id,
            "sectionName": sec.section_name,
            "availabilityScore": round((sec_window - sec_downtime) / sec_window, 4) if sec_window else 1.0,
            "scheduledDowntimeMin": sec_downtime,
            "totalWindowMin": sec_window,
        })

    return {
        "success": True,
        "data": {
            "totalTasks": total_tasks,
            "pendingTasks": pending_tasks,
            "scheduledBlocks": scheduled_blocks,
            "approvedBlocks": approved_blocks,
            "overdueCritical": overdue_critical,
            "assetAvailability": availability,
            "tasksByDept": dept_pending,
            "allTasksByDept": dept_all,
            "pendingTasksByDept": dept_pending,
            "criticalityDistribution": crit_dist,
            "availabilityTrend": trend,
            "recentSchedules": [
                {
                    "scheduleId": s.schedule_id,
                    "sectionName": s.section_name,
                    "departments": s.departments,
                    "status": s.status,
                    "windowStart": s.window_start.isoformat() if s.window_start else None,
                }
                for s in recent_scheds
            ],
            "availabilityBySection": avail_by_section,
        },
    }


@router.get("/downtime")
def downtime_report(
    plan_type: Optional[str] = Query("weekly", alias="planType"),
    zone: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    corridor_q = db.query(CorridorBlock)
    if zone and zone.strip().upper() != "ALL":
        z = zone.strip()
        corridor_q = corridor_q.filter(
            (CorridorBlock.zone_code.ilike(z)) | (CorridorBlock.zone.ilike(z))
        )
    corridors = {c.section_id: c for c in corridor_q.all()}
    horizon_minutes_map = {
        "daily": 1440,
        "weekly": 10080,
        "monthly": 43200,
    }
    horizon_minutes = horizon_minutes_map.get(plan_type, 10080)

    q = db.query(BlockSchedule)
    if zone and zone.strip().upper() != "ALL":
        z = zone.strip()
        q = q.filter(
            (BlockSchedule.zone_code.ilike(z)) |
            (BlockSchedule.section_id.in_(corridors.keys()))
        )
    if plan_type and plan_type != "all":
        q = q.filter(BlockSchedule.plan_type == plan_type)
    
    schedules = q.all()

    section_map: dict[str, dict] = {}
    for sid, c in corridors.items():
        section_map[sid] = {
            "sectionId": sid,
            "sectionName": c.section_name or sid,
            "totalDowntimeMinutes": 0,
            "approvedDowntimeMinutes": 0,
            "proposedDowntimeMinutes": 0,
            "blockCount": 0,
            "trafficDensity": c.traffic_density or "high",
            "nightMinutes": 0,
            "morningMinutes": 0,
            "middayMinutes": 0,
            "eveningMinutes": 0,
            "engineeringMinutes": 0,
            "signalTelecomMinutes": 0,
            "tractionMinutes": 0,
        }

    for s in schedules:
        sid = s.section_id
        if sid not in section_map:
            c = corridors.get(sid)
            section_map[sid] = {
                "sectionId": sid,
                "sectionName": s.section_name or (c.section_name if c else sid),
                "totalDowntimeMinutes": 0,
                "approvedDowntimeMinutes": 0,
                "proposedDowntimeMinutes": 0,
                "blockCount": 0,
                "trafficDensity": c.traffic_density if c else "high",
                "nightMinutes": 0,
                "morningMinutes": 0,
                "middayMinutes": 0,
                "eveningMinutes": 0,
                "engineeringMinutes": 0,
                "signalTelecomMinutes": 0,
                "tractionMinutes": 0,
            }
        
        dur = s.total_duration_min or 0
        section_map[sid]["totalDowntimeMinutes"] += dur
        section_map[sid]["blockCount"] += 1
        if s.status in ("approved", "executed"):
            section_map[sid]["approvedDowntimeMinutes"] += dur
        else:
            section_map[sid]["proposedDowntimeMinutes"] += dur

        # Shift classification based on start hour
        hour = s.window_start.hour if s.window_start else 1
        if 0 <= hour < 6:
            section_map[sid]["nightMinutes"] += dur
        elif 6 <= hour < 12:
            section_map[sid]["morningMinutes"] += dur
        elif 12 <= hour < 18:
            section_map[sid]["middayMinutes"] += dur
        else:
            section_map[sid]["eveningMinutes"] += dur

        # Department distribution
        depts = s.departments or []
        dur_per_dept = dur / len(depts) if depts else dur
        for d in depts:
            if "Engineering" in d:
                section_map[sid]["engineeringMinutes"] += dur_per_dept
            elif "Signal" in d:
                section_map[sid]["signalTelecomMinutes"] += dur_per_dept
            elif "Traction" in d:
                section_map[sid]["tractionMinutes"] += dur_per_dept

    report = []
    total_network_downtime_min = 0
    for v in section_map.values():
        total_min = v["totalDowntimeMinutes"]
        total_network_downtime_min += total_min
        avail = max(70.0, min(100.0, round(((horizon_minutes - total_min) / horizon_minutes) * 100, 1)))
        report.append({
            **v,
            "totalDowntimeHours": round(total_min / 60, 1),
            "nightHours": round(v["nightMinutes"] / 60, 1),
            "morningHours": round(v["morningMinutes"] / 60, 1),
            "middayHours": round(v["middayMinutes"] / 60, 1),
            "eveningHours": round(v["eveningMinutes"] / 60, 1),
            "engineeringHours": round(v["engineeringMinutes"] / 60, 1),
            "signalTelecomHours": round(v["signalTelecomMinutes"] / 60, 1),
            "tractionHours": round(v["tractionMinutes"] / 60, 1),
            "availability": avail,
        })

    # Summary metrics
    total_sections = max(1, len(report))
    avg_avail = round(sum(r["availability"] for r in report) / total_sections, 1)

    return {
        "success": True,
        "data": report,
        "meta": {
            "planType": plan_type,
            "horizonMinutes": horizon_minutes,
            "totalDowntimeMinutes": total_network_downtime_min,
            "totalDowntimeHours": round(total_network_downtime_min / 60, 1),
            "avgAvailability": avg_avail,
            "corridorCount": total_sections,
        }
    }


@router.get("/utilization")
def utilization_report(
    plan_type: Optional[str] = Query("weekly", alias="planType"),
    zone: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    q = db.query(BlockSchedule)
    if zone and zone.strip().upper() != "ALL":
        z = zone.strip()
        q = q.filter(
            (BlockSchedule.zone_code.ilike(z)) |
            (BlockSchedule.section_id.in_(
                db.query(CorridorBlock.section_id).filter(
                    (CorridorBlock.zone_code.ilike(z)) | (CorridorBlock.zone.ilike(z))
                )
            ))
        )
    if plan_type and plan_type != "all":
        q = q.filter(BlockSchedule.plan_type == plan_type)
    schedules = q.all()

    by_status: dict[str, int] = {}
    total_downtime = 0
    multi_dept = 0
    by_dept: dict[str, dict] = {
        "Engineering": {"department": "Engineering (Track)", "code": "ENG", "blockCount": 0, "totalMinutes": 0, "color": "#003366"},
        "Signal & Telecom": {"department": "Signal & Telecom", "code": "S&T", "blockCount": 0, "totalMinutes": 0, "color": "#FF671F"},
        "Traction Distribution": {"department": "Traction Distribution", "code": "TRD", "blockCount": 0, "totalMinutes": 0, "color": "#046A38"},
    }

    shift_stats = {
        "night": {"shift": "Night Block", "window": "00:00 - 06:00", "icon": "", "blockCount": 0, "totalMinutes": 0, "color": "#0B2545"},
        "morning": {"shift": "Morning Shift", "window": "06:00 - 12:00", "icon": "", "blockCount": 0, "totalMinutes": 0, "color": "#174A7E"},
        "midday": {"shift": "Midday Window", "window": "12:00 - 18:00", "icon": "", "blockCount": 0, "totalMinutes": 0, "color": "#2A75B3"},
        "evening": {"shift": "Evening Shift", "window": "18:00 - 24:00", "icon": "", "blockCount": 0, "totalMinutes": 0, "color": "#64A5D7"},
    }

    section_agg: dict[str, dict] = {}

    for s in schedules:
        st = s.status or "proposed"
        by_status[st] = by_status.get(st, 0) + 1
        dur = s.total_duration_min or 0
        total_downtime += dur
        if s.is_multi_department:
            multi_dept += 1

        # Shift
        hour = s.window_start.hour if s.window_start else 1
        if 0 <= hour < 6:
            shift_stats["night"]["blockCount"] += 1
            shift_stats["night"]["totalMinutes"] += dur
        elif 6 <= hour < 12:
            shift_stats["morning"]["blockCount"] += 1
            shift_stats["morning"]["totalMinutes"] += dur
        elif 12 <= hour < 18:
            shift_stats["midday"]["blockCount"] += 1
            shift_stats["midday"]["totalMinutes"] += dur
        else:
            shift_stats["evening"]["blockCount"] += 1
            shift_stats["evening"]["totalMinutes"] += dur

        # Departments
        depts = s.departments or []
        for d in depts:
            matched_key = None
            if "Engineering" in d:
                matched_key = "Engineering"
            elif "Signal" in d:
                matched_key = "Signal & Telecom"
            elif "Traction" in d:
                matched_key = "Traction Distribution"
            
            if matched_key:
                by_dept[matched_key]["blockCount"] += 1
                by_dept[matched_key]["totalMinutes"] += dur / len(depts)

        # Section
        sid = s.section_id
        if sid not in section_agg:
            section_agg[sid] = {"sectionId": sid, "sectionName": s.section_name or sid, "blockCount": 0, "totalMinutes": 0}
        section_agg[sid]["blockCount"] += 1
        section_agg[sid]["totalMinutes"] += dur

    total_blocks = len(schedules)
    # Add percentages and hours
    for v in shift_stats.values():
        v["totalHours"] = round(v["totalMinutes"] / 60, 1)
        v["percentage"] = round((v["blockCount"] / total_blocks * 100), 1) if total_blocks else 0

    for v in by_dept.values():
        v["totalHours"] = round(v["totalMinutes"] / 60, 1)
        v["percentage"] = round((v["blockCount"] / max(1, sum(d["blockCount"] for d in by_dept.values())) * 100), 1)

    return {
        "success": True,
        "data": {
            "totalDowntimeMinutes": total_downtime,
            "totalDowntimeHours": round(total_downtime / 60, 1),
            "totalBlocks": total_blocks,
            "byStatus": by_status,
            "multiDeptBlocks": multi_dept,
            "departmentUtilization": list(by_dept.values()),
            "shiftUtilization": list(shift_stats.values()),
            "sectionUtilization": list(section_agg.values()),
            "hoursSavedCoordinated": round(multi_dept * 1.5, 1),
        },
    }


@router.get("/audit-log")
def audit_log(
    action: Optional[str] = None,
    page: int = 1,
    limit: int = 50,
    start_date: Optional[date] = Query(None, alias="startDate"),
    end_date: Optional[date] = Query(None, alias="endDate"),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    q = db.query(AuditLog)
    if action:
        q = q.filter(AuditLog.action == action)
    if start_date:
        q = q.filter(AuditLog.created_at >= datetime.combine(start_date, datetime.min.time()))
    if end_date:
        q = q.filter(AuditLog.created_at <= datetime.combine(end_date, datetime.max.time()))

    total = q.count()
    logs = q.order_by(AuditLog.created_at.desc()).offset((page - 1) * limit).limit(limit).all()

    return {
        "success": True,
        "data": [
            {
                "id": l.id,
                "action": l.action,
                "userName": l.user_name,
                "targetId": l.target_id,
                "targetType": l.target_type,
                "details": l.details,
                "createdAt": l.created_at.isoformat() if l.created_at else None,
            }
            for l in logs
        ],
        "pagination": {"total": total, "page": page, "pages": -(-total // limit)},
    }


@router.get("/export")
def export_report(
    type: str = "schedules",
    format: str = "xlsx",
    plan_type: Optional[str] = Query(None, alias="planType"),
    department: Optional[str] = None,
    section_id: Optional[str] = Query(None, alias="sectionId"),
    status: Optional[str] = None,
    zone: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    data = []

    if type == "schedules":
        q = db.query(BlockSchedule)
        if zone and zone.strip().upper() != "ALL":
            z = zone.strip()
            q = q.filter(
                (BlockSchedule.zone_code.ilike(z)) |
                (BlockSchedule.section_id.in_(
                    db.query(CorridorBlock.section_id).filter(
                        (CorridorBlock.zone_code.ilike(z)) | (CorridorBlock.zone.ilike(z))
                    )
                ))
            )
        if plan_type and plan_type != "all":
            q = q.filter(BlockSchedule.plan_type == plan_type)
        if section_id:
            q = q.filter(BlockSchedule.section_id == section_id)
        if status:
            q = q.filter(BlockSchedule.status == status)
        rows = q.order_by(BlockSchedule.window_start).all()
        data = [
            {
                "Schedule ID": s.schedule_id,
                "Section Code": s.section_id,
                "Corridor Name": s.section_name,
                "Plan Horizon": s.plan_type.upper(),
                "Start Time": s.window_start.strftime("%d/%m/%Y %H:%M") if s.window_start else "",
                "End Time": s.window_end.strftime("%d/%m/%Y %H:%M") if s.window_end else "",
                "Duration": _fmt_duration(s.total_duration_min),
                "Departments": " + ".join(s.departments or []),
                "Multi-Dept Coordinated": "Yes" if s.is_multi_department else "No",
                "Status": (s.status or "").upper(),
                "CP-SAT Score": f"{round(s.optimizer_score * 100)}%" if s.optimizer_score else "N/A",
                "Availability Score": f"{round((s.availability_score or 0) * 100, 1)}%" if s.availability_score else "N/A",
            }
            for s in rows
        ]
    elif type == "tasks":
        q = db.query(MaintenanceTask)
        if zone and zone.strip().upper() != "ALL":
            z = zone.strip()
            q = q.filter(
                (MaintenanceTask.zone_code.ilike(z)) |
                (MaintenanceTask.section_id.in_(
                    db.query(CorridorBlock.section_id).filter(
                        (CorridorBlock.zone_code.ilike(z)) | (CorridorBlock.zone.ilike(z))
                    )
                ))
            )
        if department:
            q = q.filter(MaintenanceTask.department == department)
        if section_id:
            q = q.filter(MaintenanceTask.section_id == section_id)
        if status:
            q = q.filter(MaintenanceTask.status == status)
        rows = q.order_by(MaintenanceTask.reported_date.desc()).all()
        data = [
            {
                "Task ID": t.task_id,
                "Source System": t.source_system,
                "Department": t.department,
                "Section Code": t.section_id,
                "Corridor Name": t.section_name,
                "Defect Type": t.defect_type,
                "Criticality": (t.criticality or "").upper(),
                "Duration": _fmt_duration(t.estimated_duration),
                "Reported Date": t.reported_date.strftime("%d/%m/%Y") if t.reported_date else "",
                "Due Date": t.due_date.strftime("%d/%m/%Y") if t.due_date else "",
                "Status": (t.status or "").upper(),
                "AI Score": f"{round((t.criticality_score or 0) * 100)}%" if t.criticality_score else "N/A",
                "Urgency Tier": t.urgency_tier or "N/A",
                "Location Km": t.location_km or "",
            }
            for t in rows
        ]

    if format == "xlsx":
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Block Schedules" if type == "schedules" else "Maintenance Tasks"

        if data:
            headers = list(data[0].keys())
            ws.append(headers)
            # Style header row
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="1a3a5c", end_color="1a3a5c", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            for row in data:
                ws.append(list(row.values()))

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        filename = f"railopt_{type}_report.xlsx"
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    elif format == "pdf":
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.units import cm

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=1*cm, leftMargin=1*cm,
                                topMargin=1.5*cm, bottomMargin=1.5*cm)
        styles = getSampleStyleSheet()
        elements = []

        elements.append(Paragraph("RailOpt AI — Block Plan Report", styles["Title"]))
        elements.append(Paragraph(f"Generated: {datetime.now().strftime('%d %b %Y %H:%M IST')}", styles["Normal"]))
        elements.append(Spacer(1, 0.5*cm))

        if data:
            table_data = [list(data[0].keys())] + [list(r.values()) for r in data]
            t = Table(table_data, repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a3a5c")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f4f8")]),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cccccc")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]))
            elements.append(t)

        doc.build(elements)
        buffer.seek(0)
        filename = f"railopt_{type}_report.pdf"
        return StreamingResponse(
            buffer, media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    return {"success": True, "data": data, "count": len(data)}
